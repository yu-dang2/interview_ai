"""
이력서 라우터

모든 엔드포인트가 인증을 요구하며, 이력서는 등록한 사용자에게 귀속된다.
남의 이력서는 조회되지 않는다.

엔드포인트:
    POST /resume                           - 이력서 텍스트 직접 등록
    POST /resume/upload                    - 이력서 파일 업로드 (.pdf / .docx / .hwpx)
    GET  /resume/                          - 내 이력서 목록 조회
    GET  /resume/{resume_id}               - 특정 이력서 조회
    GET  /resume/{resume_id}/optimize      - 이력서 최적화 제안 조회
    POST /resume/{resume_id}/apply         - 최적화 제안 적용 (새 이력서 생성)
    GET  /resume/{resume_id}/download      - 이력서 docx 다운로드
"""

import io
import json

from docx import Document
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from backend.core.deps import current_user
from backend.core.masking import mask_personal_info
from backend.database import get_db
from backend.models.models import (
    InterviewResult,
    InterviewResumeOptimization,
    InterviewSession,
    Resume,
    User,
)
from backend.routers.uploads import read_document_upload
from backend.schemas.interview import ResumeOptimizationResponse, ResumeSuggestion
from backend.schemas.resume import (
    ApplyOptimizationRequest,
    ApplyOptimizationResponse,
    ResumeCreateRequest,
    ResumeListItem,
    ResumeResponse,
)

router = APIRouter(prefix="/resume", tags=["Resume"])


def _create(db: Session, content: str, user: User) -> Resume:
    # 텍스트 직접 등록 경로도 파일 업로드와 같게 비식별화한다.
    # (파일 업로드는 read_document_upload 안에서 이미 지워져서 들어온다)
    resume = Resume(content=mask_personal_info(content), users_user_id=user.user_id)
    db.add(resume)
    db.commit()
    db.refresh(resume)
    return resume


def _get_owned(db: Session, resume_id: int, user: User) -> Resume:
    """내 이력서만 반환. 없거나 남의 것이면 404.

    남의 이력서에 403 을 주면 그 id 가 존재한다는 사실이 새어나가므로 404 로 묶는다.
    """
    resume = db.get(Resume, resume_id)
    if resume is None or resume.users_user_id != user.user_id:
        raise HTTPException(status_code=404, detail="이력서를 찾을 수 없습니다.")
    return resume


@router.post("", response_model=ResumeResponse, status_code=201)
def create_resume(
    req: ResumeCreateRequest,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    """이력서 텍스트를 직접 등록한다."""
    return _create(db, req.content, user)


@router.post("/upload", response_model=ResumeResponse, status_code=201)
async def upload_resume(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    """이력서 파일(.pdf / .docx / .hwpx)을 업로드한다."""
    content = await read_document_upload(file)
    return _create(db, content, user)


@router.get("/", response_model=list[ResumeListItem])
def get_resume_list(db: Session = Depends(get_db), user: User = Depends(current_user)):
    """내 이력서 목록을 조회한다."""
    return (
        db.query(Resume)
        .filter(Resume.users_user_id == user.user_id)
        .order_by(Resume.resume_id.desc())
        .all()
    )


@router.get("/{resume_id}", response_model=ResumeResponse)
def get_resume(
    resume_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    """특정 이력서를 조회한다."""
    return _get_owned(db, resume_id, user)


# ── 이력서 최적화 ──────────────────────────────────────
# 제안 자체는 면접 종료 시 resume_optimizer 노드가 만든다. 여기서는 그 결과를
# 이력서 기준으로 꺼내 쓰고(조회), 골라서 새 이력서로 만들고(적용), 파일로 내보낸다(다운로드).

_DOWNLOAD_NOTICE = (
    "※ 개인정보 보호를 위해 이메일·연락처·주소 등은 [항목] 형태로 대체되어 있습니다. "
    "실제 지원 시에는 해당 부분을 직접 입력해주세요."
)


def _latest_optimization(db: Session, resume: Resume):
    """이 이력서로 본 면접 중 최적화 결과가 있는 가장 최근 것."""
    return (
        db.query(InterviewResumeOptimization)
        .join(InterviewResult)
        .join(InterviewSession,
              InterviewResult.interview_sessions_session_id == InterviewSession.session_id)
        .filter(InterviewSession.resumes_resume_id == resume.resume_id)
        .order_by(InterviewResumeOptimization.created_at.desc())
        .first()
    )


@router.get("/{resume_id}/optimize", response_model=ResumeOptimizationResponse)
def optimize_resume(
    resume_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    """
    이력서 최적화 제안 조회.

    이 이력서로 본 면접 중 가장 최근 결과를 돌려준다. 특정 면접의 결과가 필요하면
    GET /interview/sessions/{id}/resume-optimization 을 쓴다.
    """
    resume = _get_owned(db, resume_id, user)
    row = _latest_optimization(db, resume)
    if row is None:
        raise HTTPException(
            status_code=404,
            detail="이 이력서로 완료된 면접이 없어 최적화 제안이 없습니다.",
        )

    return ResumeOptimizationResponse(
        session_id=row.result.interview_sessions_session_id,
        resume_id=resume_id,
        matched_keywords=json.loads(row.matched_keywords or "[]"),
        missing_keywords=json.loads(row.missing_keywords or "[]"),
        suggestions=[
            ResumeSuggestion(**s)
            for s in json.loads(row.suggestions or "[]")
            if isinstance(s, dict)
        ],
    )


@router.post("/{resume_id}/apply", response_model=ApplyOptimizationResponse, status_code=201)
def apply_optimization(
    resume_id: int,
    req: ApplyOptimizationRequest,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    """
    고른 제안을 적용해 새 이력서를 만든다.

    원본을 덮어쓰지 않는 이유: 되돌릴 수 없고, 마이페이지의 이력서 버전 관리도
    원본이 남아 있어야 성립한다.
    """
    resume = _get_owned(db, resume_id, user)

    session = db.get(InterviewSession, req.session_id)
    if session is None or session.users_user_id != user.user_id:
        raise HTTPException(status_code=404, detail="면접 기록을 찾을 수 없습니다.")

    row = (
        db.query(InterviewResumeOptimization)
        .join(InterviewResult)
        .filter(InterviewResult.interview_sessions_session_id == req.session_id)
        .first()
    )
    if row is None:
        raise HTTPException(status_code=404, detail="해당 면접의 최적화 제안이 없습니다.")

    suggestions = [s for s in json.loads(row.suggestions or "[]") if isinstance(s, dict)]
    if req.suggestion_ids:
        wanted = {str(i) for i in req.suggestion_ids}
        suggestions = [s for s in suggestions if str(s.get("id")) in wanted]
    if not suggestions:
        raise HTTPException(status_code=400, detail="적용할 제안이 없습니다.")

    # original 문구를 improved 로 바꾼다. 원문에서 못 찾은 제안은 건너뛴다.
    # (LLM 이 요약해 옮겨 적은 경우 원문과 글자가 정확히 맞지 않을 수 있다)
    content = resume.content or ""
    applied = 0
    for s in suggestions:
        original = (s.get("original") or "").strip()
        improved = (s.get("improved") or "").strip()
        if original and improved and original in content:
            content = content.replace(original, improved, 1)
            applied += 1

    if applied == 0:
        raise HTTPException(
            status_code=409,
            detail="제안의 원본 문구를 이력서에서 찾지 못해 적용하지 못했습니다.",
        )

    new_resume = _create(db, content, user)
    return ApplyOptimizationResponse(
        resume_id=new_resume.resume_id,
        source_resume_id=resume_id,
        applied_count=applied,
        content=content,
    )


@router.get("/{resume_id}/download")
def download_resume(
    resume_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    """
    이력서를 docx 파일로 내려준다.

    우리가 저장하는 건 추출된 텍스트라 원본 서식은 복원되지 않는다. 최적화를 적용한
    이력서(=새로 만들어진 resume_id)를 넘기면 그 내용으로 파일이 만들어진다.

    본문에 안내 문구를 넣는 이유: 저장 시점에 연락처·주소가 비식별화되므로 파일에도
    자리표시자가 그대로 나온다. 안내가 없으면 사용자는 자기 정보가 사라진 줄 안다.
    """
    resume = _get_owned(db, resume_id, user)

    document = Document()
    document.add_paragraph(_DOWNLOAD_NOTICE)
    document.add_paragraph("")
    for line in (resume.content or "").splitlines():
        document.add_paragraph(line)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="resume_{resume_id}.docx"'
        },
    )
